from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile

import openpyxl

import YM_matcher_v2 as ym


class MatcherTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)

    def files(self, inputs, sources, header_row=1, existing=False):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Любое название"
        for _ in range(header_row - 1):
            ws.append(["Описание"])
        ws.append(["Имя артиста", "Release", "Плейлисты" if existing else "Другое"])
        for artist, release in inputs:
            ws.append([artist, release, "Старое значение"])
        ws["A1"].font = openpyxl.styles.Font(bold=True)
        wb.save(self.path / ym.INPUT_NAME)
        wb.close()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Первый лист"
        ws.append(["Дата", "Артист", "Название", "Ссылка", "E", "F", "G", "H", "I", "J", "K", "L"])
        for artist, release, playlists in sources:
            ws.append([None, artist, release, None, "+", None, None, None, playlists, None, None, "ML"])
        wb.create_sheet("Игнорируемый лист").append(["Без заголовков"])
        wb.save(self.path / ym.SOURCE_NAME)
        wb.close()

    def load(self, name):
        wb = openpyxl.load_workbook(self.path / name)
        self.addCleanup(wb.close)
        return wb.active

    def test_shared_artist_release_and_missing_report(self):
        self.files(
            [("A, B", "C"), ("A", "D"), ("A, X", "C (Live)"), ("Z", "C")],
            [("B, A", "C (Remix)", "C support"), ("A", "D", "D support"),
             ("A, Q", "E", "E support"), ("Y", "C", "Y support")],
        )
        ym.run(self.path)
        ws = self.load(ym.OUTPUT_NAME)
        self.assertEqual([ws.cell(r, 3).value for r in range(2, 6)],
                         ["C support\nML", "D support\nML", "C support\nML", None])
        self.assertEqual(ws["D2"].value, "Старое значение")
        self.assertTrue(ws["A1"].font.bold)
        missing = self.load(ym.MISSING_NAME)
        self.assertEqual([(missing.cell(r, 2).value, missing.cell(r, 3).value)
                          for r in range(2, missing.max_row + 1)], [("A, Q", "E"), ("Y", "C")])

    def test_empty_release_and_header_position_and_repeated_run(self):
        self.files([("A", None), ("A", "   "), ("A", "C")], [("A", "C", "C support")], header_row=4, existing=True)
        ym.run(self.path)
        ws = self.load(ym.OUTPUT_NAME)
        self.assertEqual(ws["C4"].value, "Плейлисты")
        self.assertIsNone(ws["C5"].value)
        self.assertIsNone(ws["C6"].value)
        self.assertEqual(ws["C7"].value, "C support\nML")
        self.assertEqual(ws.max_column, 3)
        issues = self.load(ym.ERROR_NAME)
        self.assertEqual([issues.cell(r, 3).value for r in (2, 3)], [5, 6])
        ym.run(self.path)

    def test_support_dedup_flags_and_l(self):
        self.files([("A", "C")], [("A", "C", " x На вайбе\nне поддержали\nНА ВАЙБЕ")])
        path = self.path / ym.SOURCE_NAME
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["F2"], ws["G2"], ws["H2"] = "+", "+", "х"
        ws["K2"] = "На вайбе"
        wb.save(path)
        wb.close()
        ym.run(self.path)
        self.assertEqual(self.load(ym.OUTPUT_NAME)["C2"].value, "F\nG\nH\nНа вайбе\nML")

    def test_normalization(self):
        self.assertEqual(ym.artist_tokens("A feat. B"), {"a", "b"})
        self.assertEqual(ym.release_key("  ПЕСНЯ (Live (2026)) (Remix) "), "песня")
        self.assertEqual(ym.artist_tokens("A; B & C x D"), {"a", "b", "c", "d"})

    def test_missing_report_preserves_styles_between_workbooks(self):
        self.files([("A", "C")], [("B", "D", "Support")])
        path = self.path / ym.SOURCE_NAME
        wb = openpyxl.load_workbook(path)
        wb.active["B2"].font = openpyxl.styles.Font(name="Arial", color="FF009900", bold=True)
        wb.active["B2"].alignment = openpyxl.styles.Alignment(wrap_text=True, horizontal="center")
        wb.active["B2"].number_format = "@"
        wb.save(path)
        wb.close()
        ym.run(self.path)
        cell = self.load(ym.MISSING_NAME)["B2"]
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.alignment.horizontal, "center")
        self.assertEqual(cell.number_format, "@")

    def formula_file(self, cache):
        path = self.path / "formula.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = "=1+1"
        wb.create_sheet("Второй")["B2"] = "=1+1"
        wb.save(path)
        wb.close()
        if cache:
            tmp = self.path / "patched.xlsx"
            with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
                for info in src.infolist():
                    data = src.read(info.filename)
                    if info.filename.startswith("xl/worksheets/"):
                        data = data.replace(b"<v></v>", b"<v>2</v>")
                    dst.writestr(info, data)
            tmp.replace(path)
        return path

    def test_formula_conversion_all_sheets_without_backup(self):
        path = self.formula_file(cache=True)
        self.assertEqual(ym.replace_formulas(path), 2)
        wb = openpyxl.load_workbook(path)
        try:
            self.assertEqual(wb.worksheets[0]["A1"].value, 2)
            self.assertEqual(wb.worksheets[1]["B2"].value, 2)
        finally:
            wb.close()
        self.assertEqual(list(self.path.iterdir()), [path])

    def test_missing_formula_cache_leaves_source_unchanged(self):
        path = self.formula_file(cache=False)
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "Не сохранён результат"):
            ym.replace_formulas(path)
        self.assertEqual(path.read_bytes(), before)

    def test_invalid_headers_do_not_create_results(self):
        self.files([("A", "C")], [("A", "C", "Support")])
        path = self.path / ym.INPUT_NAME
        wb = openpyxl.load_workbook(path)
        wb.active["B1"] = "Неверный заголовок"
        wb.save(path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "не найдены заголовки"):
            ym.run(self.path)
        self.assertFalse((self.path / ym.OUTPUT_NAME).exists())


if __name__ == "__main__":
    unittest.main()
